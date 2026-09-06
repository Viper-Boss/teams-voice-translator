from __future__ import annotations

import json
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path
from collections.abc import Callable

from ..aliyun import ApiError, _extract_api_error
from .memory import MeetingMemory

MAX_CONTEXT_CHARS = 12000
MAX_TERMS = 40

_CONTEXT_INSTRUCTION = """你是一名学术翻译顾问。下面是一位硕士研究生的答辩 PPT 文字内容。
请提取：
1. "brief"：用 4~8 句中文概括论文主题、研究方法和主要结论，供现场口译员理解背景。
2. "terms"：论文中的专业术语中英对照表（最多 {max_terms} 条），只收论文里真正使用的领域术语，
   不收普通词汇。这些术语在现场翻译中会被强制使用。
只输出一个 JSON 对象，格式：
{{"brief": "...", "terms": [{{"source": "中文术语", "target": "English term"}}]}}"""


@dataclass
class PptContext:
    brief: str = ""
    terms: list[dict[str, str]] = field(default_factory=list)
    slide_count: int = 0
    source_path: str = ""


def extract_pptx_text(path: Path, *, include_notes: bool = True) -> tuple[str, int]:
    """Extract all slide text (optionally speaker notes) from a .pptx file."""
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - environment guard
        raise ApiError("缺少 python-pptx，无法解析 PPT，请先安装依赖") from exc
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ApiError(f"无法打开 PPT 文件：{exc}") from exc
    chunks: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        pieces: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    pieces.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    line = " | ".join(cell for cell in cells if cell)
                    if line:
                        pieces.append(line)
        if include_notes and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                pieces.append(f"备注：{notes}")
        if pieces:
            chunks.append(f"【第{index}页】\n" + "\n".join(pieces))
    return "\n\n".join(chunks), len(presentation.slides)


def extract_plain_text(path: Path) -> tuple[str, int]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ApiError(f"无法读取文件：{exc}") from exc
    return text, 0


def extract_document_text(path: Path, *, include_notes: bool = True) -> tuple[str, int]:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return extract_pptx_text(path, include_notes=include_notes)
    if suffix in {".txt", ".md"}:
        return extract_plain_text(path)
    raise ApiError("支持 .pptx、.txt、.md 文件；如为 .ppt 请先在 PowerPoint 中另存为 .pptx")


def build_context_payload(document_text: str) -> dict:
    trimmed = document_text.strip()[:MAX_CONTEXT_CHARS]
    return {
        "model": "qwen-plus",
        "messages": [
            {"role": "system", "content": _CONTEXT_INSTRUCTION.format(max_terms=MAX_TERMS)},
            {"role": "user", "content": trimmed},
        ],
        "temperature": 0.2,
    }


def parse_context_response(content: str) -> tuple[str, list[dict[str, str]]]:
    raw = content.strip()
    fence = re.search(r"```(?:json)?\s*(.+?)```", raw, re.DOTALL)
    if fence:
        raw = fence.group(1).strip()
    start = raw.find("{")
    end = raw.rfind("}")
    if start == -1 or end == -1:
        raise ApiError("PPT 分析返回内容不是 JSON，请重试")
    try:
        data = json.loads(raw[start : end + 1])
    except json.JSONDecodeError as exc:
        raise ApiError(f"PPT 分析返回 JSON 解析失败：{exc}") from exc
    brief = str(data.get("brief", "")).strip()
    terms: list[dict[str, str]] = []
    for item in data.get("terms") or []:
        if isinstance(item, dict):
            source = str(item.get("source", "")).strip()
            target = str(item.get("target", "")).strip()
            if source and target:
                terms.append({"source": source, "target": target})
    if not brief:
        raise ApiError("PPT 分析未返回论文摘要")
    return brief, terms[:MAX_TERMS]


def import_presentation(
    client,
    path: Path,
    *,
    model: str = "qwen-plus",
    on_status: Callable[[str], None] | None = None,
    cancel_event: threading.Event | None = None,
    include_notes: bool = True,
) -> PptContext:
    """Read a PPT/file, then distill a defense brief + glossary via one LLM call."""
    notify = on_status or (lambda message: None)
    notify("正在读取文件…")
    document_text, slide_count = extract_document_text(Path(path), include_notes=include_notes)
    if not document_text.strip():
        raise ApiError("文件中没有可提取的文字")
    if cancel_event is not None and cancel_event.is_set():
        raise ApiError("已取消")
    notify(f"正在通读全文并提炼术语（{len(document_text)} 字）…")
    payload = build_context_payload(document_text)
    payload["model"] = model or payload["model"]
    import requests

    if hasattr(client, "endpoint"):  # ContextTranslator
        endpoint = client.endpoint
    else:  # BailianClient
        endpoint = f"{client.root}/compatible-mode/v1/chat/completions"
    response = requests.post(
        endpoint,
        headers={"Authorization": f"Bearer {client.api_key}", "Content-Type": "application/json"},
        json=payload,
        timeout=max(getattr(client, "timeout", 45), 120),
        proxies=getattr(client, "proxies", None),
    )
    if cancel_event is not None and cancel_event.is_set():
        raise ApiError("已取消")
    if not response.ok:
        raise ApiError(_extract_api_error(response))
    try:
        content = str(response.json()["choices"][0]["message"]["content"])
    except Exception as exc:
        raise ApiError(f"PPT 分析返回格式异常：{response.text[:300]}") from exc
    brief, terms = parse_context_response(content)
    return PptContext(
        brief=brief,
        terms=terms,
        slide_count=slide_count,
        source_path=str(path),
    )


def apply_to_memory(memory: MeetingMemory, context: PptContext) -> None:
    memory.set_brief(context.brief)
    memory.set_glossary(context.terms)
